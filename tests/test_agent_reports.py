import pytest
from io import BytesIO
from openpyxl import load_workbook
from scarascope.core.config import Settings
from scarascope.core.models import ResearchRequest
from scarascope.services.agent import research_company
from scarascope.reports.pdf import export_pdf
from scarascope.reports.excel import export_excel, safe_cell

@pytest.mark.asyncio
async def test_offline_agent_has_bounded_plan_and_verified_demo():
    report=await research_company(ResearchRequest(ticker="DEMO",provider="none"),Settings())
    assert report.research.data_kind=="synthetic_demo"
    assert report.verification.passed
    assert len(report.plan.tasks)<=report.plan.limit
    assert len(report.steps)==len(report.plan.tasks)
    assert report.filing.evidence==[]
    assert "fictional" in report.executive_summary.lower()
    assert "No retrieved narrative" in " ".join(report.limitations)

@pytest.mark.asyncio
async def test_pdf_and_excel_exports_are_readable_and_source_aware():
    report=await research_company(ResearchRequest(ticker="DEMO",provider="none"),Settings())
    pdf=export_pdf(report)
    xlsx=export_excel(report)
    assert pdf.startswith(b"%PDF-")
    wb=load_workbook(BytesIO(xlsx))
    assert {"Financials","SEC Sources","Narrative Evidence","Caveats"}.issubset(wb.sheetnames)
    assert wb["Financials"]["A2"].value
    assert wb["Financials"]["F2"].value.startswith("=IFERROR")
    assert wb["Caveats"]["B1"].value.startswith("Synthetic")
    assert safe_cell("=1+1")=="'=1+1"
